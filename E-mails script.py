import os
from openpyxl import Workbook, load_workbook, workbook
Seguros = ["Multi Bônus", "Bolsa Premiada", "Casa tranquila", "Assistência Saúde", "Auto e Moto", "Assistência Mulher", "Pet Básico", "Pet Completo", "Odonto", "Odonto Plus"]
genero = ["M", "F"]
def email_base():
    if genero_cliente == 0:
        email_cliente = """Victor, boa tarde!

Tudo bem? 

Pode verificar por gentileza
Cliente entrou em contato com o SAC solicitando estorno do seguro {}.
O mesmo está ciente de que possui termo de adesão assinado, mas informa não ter ciência da cobrança.
Temos como verificar uma possibilidade de estorno do valor gerado por conta do seguro.

Cliente: {}
CPF: {}

""".format(Seguros[seguro_cliente], nome_cliente, cpf_cliente)
    elif genero_cliente == 1:
        email_cliente = """Victor, boa tarde!

Tudo bem? 

Pode verificar por gentileza
Cliente entrou em contato com o SAC solicitando estorno do seguro {}.
A mesma está ciente de que possui termo de adesão assinado, mas informa não ter ciência da cobrança.
Temos como verificar uma possibilidade de estorno do valor gerado por conta do seguro.

Cliente: {}
CPF: {}
""".format(Seguros[seguro_cliente], nome_cliente, cpf_cliente)
    return email_cliente
def criar_planilha():
    email_montado = "emails_montados.xlsx"
    if os.path.exists(email_montado):
        planilha = load_workbook(email_montado)
    else:
        planilha = Workbook()
        aba = planilha.active
        aba.title = "emails"
        aba["A1"] = "Nome"
        aba["B1"] = "CPF"
        aba["C1"] = "E-mail"
        planilha.save(email_montado)
        print("Planilha criada!")
def salvar_planilha(nome_cliente, cpf_cliente, email_cliente):
    email_montado = "emails_montados.xlsx"
    planilha = load_workbook(email_montado)
    aba = planilha.active
    nova_linha = [nome_cliente, cpf_cliente, email_cliente]
    aba.append(nova_linha)
    planilha.save(email_montado)
    print("E-mail salvo na planilha!")
while True:
    criar_planilha()
    nome_cliente = str(input("Digite o nome do cliente: ").strip())
    cpf_cliente = str(input("Digite o CPF do cliente: ").strip())
    genero_cliente = int(input("Qual o genero do cliente:\n [0] Masculino\n [1] Feminino\nDigite a opção: "))
    seguro_cliente = int(input("Qual o seguro do cliente:\n [0] Multi Bônus\n [1] Bolsa Premiada\n [2] Casa tranquila\n [3] Assistência Saúde\n [4] Auto e Moto\n [5] Assistência Mulher\n [6] Pet Básico\n [7] Pet Completo\n [8] Odonto\n [9] Odonto Plus\nDigite a opção: "))
    email_cliente = email_base()
    salvar_planilha(nome_cliente, cpf_cliente, email_cliente)
